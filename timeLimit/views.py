from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Count, Q
from django.db.models.functions import TruncMonth
from django.http import HttpResponse, FileResponse, Http404
from django.contrib import messages

from .models import TimeLimit
from .forms import TimeLimitForm
from .google_sheet_sync import (
    save_record_to_sheet,
    delete_record_from_sheet,
    sync_sheet_to_portal,
)
from .google_drive import (
    upload_file_to_drive,
    TL_PDF_FOLDER_ID,
    ANSWER_PDF_FOLDER_ID,
)

import openpyxl
from openpyxl import Workbook

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from html import escape
from datetime import datetime


# ============================================================
# DASHBOARD
# ============================================================

def dashboard(request):

    total_tl = TimeLimit.objects.count()

    pending_tl = TimeLimit.objects.filter(
        current_status='Pending'
    ).count()

    disposed_tl = TimeLimit.objects.filter(
        current_status='Disposed'
    ).count()

    total_sections = (
        TimeLimit.objects
        .values('section_name')
        .distinct()
        .count()
    )

    # --------------------------------------------------------
    # PDF COUNTS
    # --------------------------------------------------------

    tl_with_pdf = (
        TimeLimit.objects
        .filter(tl_pdf__isnull=False)
        .exclude(tl_pdf='')
        .count()
    )

    tl_without_pdf = total_tl - tl_with_pdf

    answer_with_pdf = (
        TimeLimit.objects
        .filter(answer_pdf__isnull=False)
        .exclude(answer_pdf='')
        .count()
    )

    # --------------------------------------------------------
    # SECTION-WISE TOTAL
    # --------------------------------------------------------

    section_data = list(
        TimeLimit.objects
        .values('section_name')
        .annotate(
            total=Count('id')
        )
        .order_by('-total')
    )

    max_count = (
        section_data[0]['total']
        if section_data
        else 1
    )

    for item in section_data:

        item['section_name'] = (
            item['section_name']
            or 'Unassigned'
        )

        item['percentage'] = round(
            (item['total'] / max_count) * 100
        )

    # --------------------------------------------------------
    # SECTION-WISE PENDING / DISPOSED
    # --------------------------------------------------------

    section_status_data = []

    sections = (
        TimeLimit.objects
        .values_list(
            'section_name',
            flat=True
        )
        .distinct()
        .order_by('section_name')
    )

    for section in sections:

        section_name = section or 'Unassigned'

        if section:

            section_records = TimeLimit.objects.filter(
                section_name=section
            )

        else:

            section_records = (
                TimeLimit.objects.filter(
                    section_name__isnull=True
                )
                |
                TimeLimit.objects.filter(
                    section_name=''
                )
            )

        total = section_records.count()

        pending = section_records.filter(
            current_status='Pending'
        ).count()

        disposed = section_records.filter(
            current_status='Disposed'
        ).count()

        section_status_data.append({
            'section_name': section_name,
            'total': total,
            'pending': pending,
            'disposed': disposed,
        })

    # --------------------------------------------------------
    # MONTH-WISE TL
    # --------------------------------------------------------

    monthly_queryset = (
        TimeLimit.objects
        .exclude(received_date__isnull=True)
        .annotate(
            month=TruncMonth('received_date')
        )
        .values('month')
        .annotate(
            total=Count('id')
        )
        .order_by('month')
    )

    monthly_data = []

    max_month_count = 1

    for item in monthly_queryset:

        if item['month']:

            month_name = item['month'].strftime(
                '%b %Y'
            )

            monthly_data.append({
                'month': month_name,
                'total': item['total'],
            })

            if item['total'] > max_month_count:
                max_month_count = item['total']

    for item in monthly_data:

        item['percentage'] = round(
            (item['total'] / max_month_count) * 100
        )

    # --------------------------------------------------------
    # RECENT RECORDS
    # --------------------------------------------------------

    recent_records = (
        TimeLimit.objects
        .order_by('-created_at')[:8]
    )

    # --------------------------------------------------------
    # PERCENTAGE
    # --------------------------------------------------------

    pending_percentage = (
        round((pending_tl / total_tl) * 100)
        if total_tl
        else 0
    )

    disposed_percentage = (
        round((disposed_tl / total_tl) * 100)
        if total_tl
        else 0
    )

    # --------------------------------------------------------
    # CONTEXT
    # --------------------------------------------------------

    context = {

        'total_tl': total_tl,

        'pending_tl': pending_tl,

        'disposed_tl': disposed_tl,

        'total_sections': total_sections,

        'tl_with_pdf': tl_with_pdf,

        'tl_without_pdf': tl_without_pdf,

        'answer_with_pdf': answer_with_pdf,

        'pending_percentage': pending_percentage,

        'disposed_percentage': disposed_percentage,

        'section_data': section_data,

        'section_status_data': section_status_data,

        'monthly_data': monthly_data,

        'recent_records': recent_records,

    }

    return render(
        request,
        'timeLimit/dashboard.html',
        context
    )


# ============================================================
# TL RECORDS
# ============================================================

def tl_records(request):

    try:

        sync_sheet_to_portal()

    except Exception as e:

        messages.warning(
            request,
            f'Google Sheet sync failed: {e}'
        )

    records = TimeLimit.objects.all()

    # --------------------------------------------------------
    # SEARCH
    # --------------------------------------------------------

    search = request.GET.get(
        'search',
        ''
    ).strip()

    if search:

        records = records.filter(

            Q(tl_no__icontains=search) |

            Q(subject__icontains=search) |

            Q(sender_name__icontains=search) |

            Q(section_name__icontains=search) |

            Q(description__icontains=search) |

            Q(current_situation__icontains=search) |

            Q(letter_no_date__icontains=search)

        )

    # --------------------------------------------------------
    # S.NO FILTER
    # --------------------------------------------------------

    sno = request.GET.get(
        'sno',
        ''
    ).strip()

    if sno:

        try:

            records = records.filter(
                sno=int(sno)
            )

        except ValueError:

            pass

    # --------------------------------------------------------
    # SECTION FILTER
    # --------------------------------------------------------

    section = request.GET.get(
        'section',
        ''
    ).strip()

    if section:

        records = records.filter(
            section_name=section
        )

    # --------------------------------------------------------
    # YEAR FILTER
    # --------------------------------------------------------

    year = request.GET.get(
        'year',
        ''
    ).strip()

    if year:

        try:

            records = records.filter(
                received_date__year=int(year)
            )

        except ValueError:

            pass

    # --------------------------------------------------------
    # STATUS FILTER
    # --------------------------------------------------------

    status = request.GET.get(
        'status',
        ''
    ).strip()

    if status:

        records = records.filter(
            current_status=status
        )

    # --------------------------------------------------------
    # ORDER
    # --------------------------------------------------------

    records = records.order_by(
        'sno',
        'id'
    )

    # --------------------------------------------------------
    # SECTION LIST
    # --------------------------------------------------------

    sections = (
        TimeLimit.objects
        .values_list(
            'section_name',
            flat=True
        )
        .distinct()
        .order_by('section_name')
    )

    # --------------------------------------------------------
    # YEAR LIST
    # --------------------------------------------------------

    years = (
        TimeLimit.objects
        .dates(
            'received_date',
            'year',
            order='DESC'
        )
    )

    context = {

        'records': records,

        'sections': sections,

        'years': years,

        'search': search,

        'sno': sno,

        'selected_section': section,

        'selected_year': year,

        'selected_status': status,

    }

    return render(
        request,
        'timeLimit/tl_records.html',
        context
    )


# ============================================================
# TL DETAIL
# ============================================================

def tl_detail(request, pk):

    records = list(
        TimeLimit.objects
        .order_by(
            'sno',
            'id'
        )
    )

    current_index = next(
        (
            index
            for index, record
            in enumerate(records)
            if record.pk == pk
        ),
        None
    )

    if current_index is None:

        return HttpResponse(
            'TL record not found.',
            status=404
        )

    record = records[current_index]

    # --------------------------------------------------------
    # PREVIOUS
    # --------------------------------------------------------

    if current_index > 0:

        previous_record = records[
            current_index - 1
        ]

    else:

        previous_record = None

    # --------------------------------------------------------
    # NEXT
    # --------------------------------------------------------

    if current_index < len(records) - 1:

        next_record = records[
            current_index + 1
        ]

    else:

        next_record = None

    return render(
        request,
        'timeLimit/tl_detail.html',
        {

            'record': record,

            'previous_record': previous_record,

            'next_record': next_record,

            'is_first': previous_record is None,

            'is_last': next_record is None,

            'current_number': current_index + 1,

            'total_records': len(records),

        }
    )


# ============================================================
# THANK YOU
# ============================================================

def thank_you(request):

    return render(
        request,
        'timeLimit/thank_you.html'
    )


# ============================================================
# ADD TL
# ============================================================

def add_tl(request):

    if request.method == 'POST':

        form = TimeLimitForm(
            request.POST,
            request.FILES
        )

        print("POST:", request.POST)
        print("FILES:", request.FILES)
        print("TL PDF:", request.FILES.get("tl_pdf"))
        print("ANSWER PDF:", request.FILES.get("answer_pdf"))

        if form.is_valid():

            record = form.save()

            try:

                # =========================
                # TL PDF → GOOGLE DRIVE
                # =========================

                if record.tl_pdf:

                    with open(record.tl_pdf.path, "rb") as pdf_file:

                        uploaded_file = upload_file_to_drive(
                            pdf_file,
                            record.tl_pdf.name.split("/")[-1],
                            TL_PDF_FOLDER_ID
                        )

                    record.tl_pdf_drive_url = (
                        uploaded_file.get("webViewLink")
                        or uploaded_file.get("webContentLink")
                        or ""
                    )


                # =========================
                # ANSWER PDF → GOOGLE DRIVE
                # =========================

                if record.answer_pdf:

                    with open(record.answer_pdf.path, "rb") as pdf_file:

                        uploaded_file = upload_file_to_drive(
                            pdf_file,
                            record.answer_pdf.name.split("/")[-1],
                            ANSWER_PDF_FOLDER_ID
                        )

                    record.answer_pdf_drive_url = (
                        uploaded_file.get("webViewLink")
                        or uploaded_file.get("webContentLink")
                        or ""
                    )


                # =========================
                # SAVE DRIVE LINKS
                # =========================

                record.save()


                # =========================
                # GOOGLE SHEET
                # =========================

                save_record_to_sheet(record)


                messages.success(
                    request,
                    'TL Record successfully added, PDFs uploaded to Google Drive and Google Sheet updated.'
                )


            except Exception as e:

                print(
                    "GOOGLE DRIVE/SHEET ERROR:",
                    repr(e)
                )

                messages.warning(
                    request,
                    f'TL added, but Google Drive/Sheet sync failed: {e}'
                )


            return redirect('tl_records')


    else:

        form = TimeLimitForm()


    return render(
        request,
        'timeLimit/tl_form.html',
        {
            'form': form,
            'title': 'Add New TL'
        }
    )
# ============================================================
# EDIT TL
# ============================================================

def edit_tl(request, pk):

    record = get_object_or_404(
        TimeLimit,
        pk=pk
    )

    if request.method == 'POST':

        form = TimeLimitForm(
            request.POST,
            request.FILES,
            instance=record
        )

        if form.is_valid():

            record = form.save()

            try:

                # -----------------------------
                # Upload NEW TL PDF to Google Drive
                # -----------------------------
                if 'tl_pdf' in request.FILES:

                    with record.tl_pdf.open('rb') as pdf_file:

                        uploaded_file = upload_file_to_drive(
                            pdf_file,
                            record.tl_pdf.name.split('/')[-1],
                            TL_PDF_FOLDER_ID
                        )

                    record.tl_pdf_drive_url = (
                        uploaded_file.get('webViewLink')
                        or uploaded_file.get('webContentLink')
                        or ''
                    )

                # -----------------------------
                # Upload NEW Answer PDF to Google Drive
                # -----------------------------
                if 'answer_pdf' in request.FILES:

                    with record.answer_pdf.open('rb') as pdf_file:

                        uploaded_file = upload_file_to_drive(
                            pdf_file,
                            record.answer_pdf.name.split('/')[-1],
                            ANSWER_PDF_FOLDER_ID
                        )

                    record.answer_pdf_drive_url = (
                        uploaded_file.get('webViewLink')
                        or uploaded_file.get('webContentLink')
                        or ''
                    )

                record.save()

                # -----------------------------
                # Update Google Sheet
                # -----------------------------
                save_record_to_sheet(record)

                messages.success(
                    request,
                    'TL Record successfully updated, PDFs synced to Google Drive and Google Sheet updated.'
                )

            except Exception as e:

                messages.warning(
                    request,
                    f'TL updated, but Google Drive/Sheet sync failed: {e}'
                )

            return redirect(
                'tl_records'
            )

    else:

        form = TimeLimitForm(
            instance=record
        )

    return render(
        request,
        'timeLimit/tl_form.html',
        {
            'form': form,
            'title': 'Edit TL'
        }
    )


# ============================================================
# DELETE TL
# ============================================================

def delete_tl(request, pk):

    record = get_object_or_404(
        TimeLimit,
        pk=pk
    )

    if request.method == 'POST':

        tl_no = record.tl_no

        record.delete()

        try:

            delete_record_from_sheet(tl_no)

            messages.success(
                request,
                'TL Record deleted successfully from Portal and Google Sheet.'
            )

        except Exception as e:

            messages.warning(
                request,
                f'TL deleted from Portal, but Google Sheet sync failed: {e}'
            )

        return redirect(
            'tl_records'
        )

    return redirect(
        'tl_records'
    )


# ============================================================
# REPORTS
# ============================================================

def reports(request):

    records = TimeLimit.objects.all()

    status = request.GET.get(
        'status',
        ''
    )

    section = request.GET.get(
        'section',
        ''
    )

    year = request.GET.get(
        'year',
        ''
    )

    if status:

        records = records.filter(
            current_status=status
        )

    if section:

        records = records.filter(
            section_name=section
        )

    if year:

        try:

            records = records.filter(
                received_date__year=int(year)
            )

        except ValueError:

            pass

    records = records.order_by(
        'sno',
        'id'
    )

    sections = (
        TimeLimit.objects
        .values_list(
            'section_name',
            flat=True
        )
        .distinct()
        .order_by(
            'section_name'
        )
    )

    years = (
        TimeLimit.objects
        .dates(
            'received_date',
            'year',
            order='DESC'
        )
    )

    return render(
        request,
        'timeLimit/reports.html',
        {

            'records': records,

            'sections': sections,

            'years': years,

            'selected_status': status,

            'selected_section': section,

            'selected_year': year,

        }
    )


# ============================================================
# EXCEL DOWNLOAD
# ============================================================

def excel_download(request):

    records = (
        TimeLimit.objects
        .all()
        .order_by(
            'sno',
            'id'
        )
    )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = 'TL Records'

    # --------------------------------------------------------
    # HEADERS
    # --------------------------------------------------------

    headers = [

        'स.क्र.',

        'TL No.',

        'प्राप्ति दिनांक',

        'प्रेषक',

        'क्र/दिनांक',

        'विषय',

        'संबंधित शाखा का नाम',

        'विवरण',

        'वर्तमान स्थिति',

        'स्थिति',

        'TL PDF',

        'Answer PDF',

    ]

    sheet.append(headers)

    # --------------------------------------------------------
    # RECORDS
    # --------------------------------------------------------

    for record in records:

        received_date = (

            record.received_date

            if record.received_date

            else ''

        )

        tl_pdf = (

            record.tl_pdf.url

            if record.tl_pdf

            else ''

        )

        answer_pdf = (

            record.answer_pdf.url

            if record.answer_pdf

            else ''

        )

        sheet.append([

            record.sno,

            record.tl_no,

            received_date,

            record.sender_name,

            record.letter_no_date,

            record.subject,

            record.section_name,

            record.description,

            record.current_situation,

            record.current_status,

            tl_pdf,

            answer_pdf,

        ])

    # --------------------------------------------------------
    # HEADER STYLE
    # --------------------------------------------------------

    for cell in sheet[1]:

        cell.font = cell.font.copy(
            bold=True
        )

        cell.alignment = cell.alignment.copy(

            horizontal='center',

            vertical='center',

            wrap_text=True

        )

    # --------------------------------------------------------
    # ALL CELLS
    # --------------------------------------------------------

    for row in sheet.iter_rows():

        for cell in row:

            cell.alignment = cell.alignment.copy(

                vertical='top',

                wrap_text=True

            )

    # --------------------------------------------------------
    # COLUMN WIDTH
    # --------------------------------------------------------

    widths = {

        'A': 10,

        'B': 18,

        'C': 16,

        'D': 25,

        'E': 22,

        'F': 30,

        'G': 28,

        'H': 45,

        'I': 40,

        'J': 18,

        'K': 35,

        'L': 35,

    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------------

    sheet.freeze_panes = 'A2'

    # --------------------------------------------------------
    # RESPONSE
    # --------------------------------------------------------

    response = HttpResponse(

        content_type=
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    )

    response['Content-Disposition'] = (
        'attachment; filename="TL_Records.xlsx"'
    )

    workbook.save(response)

    return response


# ============================================================
# EXCEL UPLOAD
# ============================================================

def excel_upload(request):

    if request.method == 'POST':

        excel_file = request.FILES.get('excel_file')

        if not excel_file:
            messages.error(
                request,
                'Please select an Excel file.'
            )
            return redirect('excel_upload')

        try:

            workbook = openpyxl.load_workbook(
                excel_file,
                data_only=True
            )

            sheet = workbook.active

            count = 0
            sheet_errors = []

            for row in sheet.iter_rows(
                min_row=2,
                values_only=True
            ):

                if not row:
                    continue

                row = list(row)

                while len(row) < 12:
                    row.append(None)

                sno = row[0]
                tl_no = row[1]
                received_date = row[2]
                sender_name = row[3]
                letter_no_date = row[4]
                subject = row[5]
                section_name = row[6]
                description = row[7]
                current_situation = row[8]
                current_status = row[9]

                if not tl_no:
                    continue

                try:
                    sno = int(sno)
                except (TypeError, ValueError):
                    continue

                if isinstance(received_date, datetime):
                    received_date = received_date.date()

                current_status = str(
                    current_status or ''
                ).strip()

                if current_status not in [
                    'Pending',
                    'Disposed'
                ]:
                    current_status = 'Pending'

                record = TimeLimit.objects.create(
                    sno=sno,
                    tl_no=str(tl_no).strip(),
                    received_date=received_date,
                    sender_name=str(
                        sender_name or ''
                    ).strip(),
                    letter_no_date=str(
                        letter_no_date or ''
                    ).strip(),
                    subject=str(
                        subject or ''
                    ).strip(),
                    section_name=str(
                        section_name or ''
                    ).strip(),
                    description=str(
                        description or ''
                    ).strip(),
                    current_situation=str(
                        current_situation or ''
                    ).strip(),
                    current_status=current_status,
                )

                try:
                    save_record_to_sheet(record)
                except Exception as e:
                    sheet_errors.append(
                        f'{record.tl_no}: {e}'
                    )

                count += 1

            messages.success(
                request,
                f'{count} TL records imported successfully.'
            )

            if sheet_errors:
                messages.warning(
                    request,
                    f'{len(sheet_errors)} records could not be synced to Google Sheet.'
                )

        except Exception as e:

            messages.error(
                request,
                f'Excel upload error: {e}'
            )

        return redirect('tl_records')

    return render(
        request,
        'timeLimit/excel_upload.html'
    )

# ============================================================
# PDF FONTS - MANGAL + ARIAL UNICODE
# ============================================================

def get_pdf_fonts():

    fonts = {}

    # --------------------------------------------------------
    # MANGAL
    # --------------------------------------------------------

    mangal_paths = [

        r"C:\Windows\Fonts\mangal.ttf",

        r"C:\Windows\Fonts\MANGAL.TTF",

    ]

    mangal_loaded = False

    for mangal_path in mangal_paths:

        try:

            if 'Mangal' not in pdfmetrics.getRegisteredFontNames():

                pdfmetrics.registerFont(
                    TTFont(
                        'Mangal',
                        mangal_path
                    )
                )

            fonts['main'] = 'Mangal'

            mangal_loaded = True

            break

        except Exception:

            continue

    if not mangal_loaded:

        fonts['main'] = 'Helvetica'

    # --------------------------------------------------------
    # ARIAL UNICODE MS
    # --------------------------------------------------------

    arial_unicode_paths = [

        r"C:\Windows\Fonts\ARIALUNI.TTF",

        r"C:\Windows\Fonts\arialuni.ttf",

        r"C:\Windows\Fonts\Arial Unicode MS.ttf",

    ]

    arial_loaded = False

    for arial_path in arial_unicode_paths:

        try:

            if (
                'ArialUnicode'
                not in pdfmetrics.getRegisteredFontNames()
            ):

                pdfmetrics.registerFont(
                    TTFont(
                        'ArialUnicode',
                        arial_path
                    )
                )

            fonts['unicode'] = 'ArialUnicode'

            arial_loaded = True

            break

        except Exception:

            continue

    if not arial_loaded:

        fonts['unicode'] = fonts['main']

    return fonts


# ============================================================
# PDF PAGE NUMBER
# ============================================================

def add_page_number(
    canvas,
    document
):

    canvas.saveState()

    canvas.setFont(
        'Helvetica',
        8
    )

    canvas.drawCentredString(

        landscape(A4)[0] / 2,

        12,

        f'Page {document.page}'

    )

    canvas.restoreState()


# ============================================================
# PDF DOWNLOAD
# ============================================================

def pdf_download(request):

    records = (
        TimeLimit.objects
        .all()
        .order_by(
            'sno',
            'id'
        )
    )

    response = HttpResponse(
        content_type='application/pdf'
    )

    response['Content-Disposition'] = (
        'attachment; filename="TL_Records.pdf"'
    )

    document = SimpleDocTemplate(

        response,

        pagesize=landscape(A4),

        rightMargin=18,

        leftMargin=18,

        topMargin=25,

        bottomMargin=25,

        title='TL Management Portal - TL Records',

        author='TL Management Portal'

    )

    styles = getSampleStyleSheet()

    # --------------------------------------------------------
    # FONTS
    # --------------------------------------------------------

    fonts = get_pdf_fonts()

    font_name = fonts['main']

    # --------------------------------------------------------
    # TITLE STYLE
    # --------------------------------------------------------

    title_style = ParagraphStyle(

        'MangalTitle',

        parent=styles['Title'],

        fontName=font_name,

        fontSize=15,

        leading=19,

        alignment=1,

        spaceAfter=6,

    )

    # --------------------------------------------------------
    # SUBTITLE STYLE
    # --------------------------------------------------------

    subtitle_style = ParagraphStyle(

        'MangalSubtitle',

        parent=styles['Normal'],

        fontName=font_name,

        fontSize=8,

        leading=10,

        alignment=1,

        textColor=colors.HexColor(
            '#64748b'
        ),

        spaceAfter=10,

    )

    # --------------------------------------------------------
    # BODY STYLE
    # --------------------------------------------------------

    body_style = ParagraphStyle(

        'MangalBody',

        parent=styles['Normal'],

        fontName=font_name,

        fontSize=6.4,

        leading=8,

        wordWrap='CJK',

        alignment=1,

    )

    # --------------------------------------------------------
    # HEADER STYLE
    # --------------------------------------------------------

    header_style = ParagraphStyle(

        'MangalHeader',

        parent=styles['Normal'],

        fontName=font_name,

        fontSize=6.3,

        leading=7.5,

        textColor=colors.white,

        alignment=1,

        wordWrap='CJK',

    )

    # ========================================================
    # HEADERS
    # ========================================================

    headers = [

        'स.क्र.',

        'TL No.',

        'प्राप्ति दिनांक',

        'प्रेषक',

        'क्र/दिनांक',

        'विषय',

        'संबंधित शाखा का नाम',

        'विवरण',

        'वर्तमान स्थिति',

        'स्थिति',

    ]

    table_data = []

    # --------------------------------------------------------
    # HEADER ROW
    # --------------------------------------------------------

    table_data.append([

        Paragraph(

            f'<b>{escape(str(header))}</b>',

            header_style

        )

        for header in headers

    ])

    # ========================================================
    # RECORD DATA
    # ========================================================

    for record in records:

        received_date = (

            record.received_date.strftime(
                '%d-%m-%Y'
            )

            if record.received_date

            else ''

        )

        values = [

            str(
                record.sno or ''
            ),

            str(
                record.tl_no or ''
            ),

            received_date,

            str(
                record.sender_name or ''
            ),

            str(
                record.letter_no_date or ''
            ),

            str(
                record.subject or ''
            ),

            str(
                record.section_name or ''
            ),

            str(
                record.description or ''
            ),

            # -----------------------------------------------
            # वर्तमान स्थिति
            # -----------------------------------------------

            str(
                record.current_situation or ''
            ),

            # -----------------------------------------------
            # स्थिति
            # -----------------------------------------------

            str(
                record.current_status or ''
            ),

        ]

        row_data = []

        for value in values:

            safe_value = escape(
                value
            ).replace(
                '\n',
                '<br/>'
            )

            row_data.append(

                Paragraph(

                    safe_value,

                    body_style

                )

            )

        table_data.append(
            row_data
        )

    # ========================================================
    # COLUMN WIDTHS
    # ========================================================

    col_widths = [

        32,     # S.No

        58,     # TL No.

        68,     # Received Date

        82,     # Sender

        68,     # Kr/Date

        90,     # Subject

        88,     # Section

        135,    # Description

        130,    # Current Situation

        54,     # Status

    ]

    # ========================================================
    # TABLE
    # ========================================================

    table = Table(

        table_data,

        colWidths=col_widths,

        repeatRows=1,

        splitByRow=1,

        hAlign='CENTER'

    )

    table.setStyle(

        TableStyle([

            (
                'BACKGROUND',

                (0, 0),

                (-1, 0),

                colors.HexColor(
                    '#173f5f'
                )

            ),

            (
                'TEXTCOLOR',

                (0, 0),

                (-1, 0),

                colors.white

            ),

            (
                'GRID',

                (0, 0),

                (-1, -1),

                0.35,

                colors.HexColor(
                    '#94a3b8'
                )

            ),

            (
                'ALIGN',

                (0, 0),

                (-1, -1),

                'CENTER'

            ),

            (
                'VALIGN',

                (0, 0),

                (-1, -1),

                'MIDDLE'

            ),

            (
                'LEFTPADDING',

                (0, 0),

                (-1, -1),

                3

            ),

            (
                'RIGHTPADDING',

                (0, 0),

                (-1, -1),

                3

            ),

            (
                'TOPPADDING',

                (0, 0),

                (-1, -1),

                4

            ),

            (
                'BOTTOMPADDING',

                (0, 0),

                (-1, -1),

                4

            ),

            (
                'ROWBACKGROUNDS',

                (0, 1),

                (-1, -1),

                [

                    colors.white,

                    colors.HexColor(
                        '#f8fafc'
                    )

                ]

            ),

        ])

    )

    # ========================================================
    # STORY
    # ========================================================

    story = [

        Paragraph(
            'TL Management Portal',
            title_style
        ),

        Paragraph(
            'TL Records Report',
            subtitle_style
        ),

        table,

        Spacer(
            1,
            8
        ),

    ]

    # ========================================================
    # BUILD
    # ========================================================

    document.build(

        story,

        onFirstPage=add_page_number,

        onLaterPages=add_page_number

    )

    return response


# ============================================================
# SERVE PDF / PDF PREVIEW
# ============================================================

def serve_pdf(
    request,
    pk,
    file_type='tl'
):

    record = get_object_or_404(
        TimeLimit,
        pk=pk
    )

    # --------------------------------------------------------
    # TL PDF
    # --------------------------------------------------------

    if file_type == 'tl':

        pdf_file = record.tl_pdf

    # --------------------------------------------------------
    # ANSWER PDF
    # --------------------------------------------------------

    elif file_type == 'answer':

        pdf_file = record.answer_pdf

    # --------------------------------------------------------
    # INVALID TYPE
    # --------------------------------------------------------

    else:

        raise Http404(
            'Invalid PDF type'
        )

    # --------------------------------------------------------
    # FILE NOT FOUND
    # --------------------------------------------------------

    if not pdf_file:

        raise Http404(
            'PDF not found'
        )

    # --------------------------------------------------------
    # OPEN PDF
    # --------------------------------------------------------

    try:

        pdf_file.open('rb')

        response = FileResponse(

            pdf_file,

            content_type='application/pdf'

        )

        response[
            'Content-Disposition'
        ] = 'inline'

        response[
            'X-Frame-Options'
        ] = 'SAMEORIGIN'

        return response

    except Exception:

        raise Http404(
            'Unable to open PDF'
        )